from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
root=Path('D:/PRM393/PRM.POWER9');build=root/'.build/presentation/office-2026';out=root/'output/presentation/Power9_Office_Lab2'
narration=json.loads((build/'narration.json').read_text(encoding='utf8'));slides=json.loads((build/'speaker-notes.json').read_text(encoding='utf8'))
fig='https://www.figma.com/design/MTv5qJbn065VYJQIXausoq';stitch='https://stitch.withgoogle.com/projects/2176700449127941467'
def proto(node):return 'https://www.figma.com/proto/MTv5qJbn065VYJQIXausoq?node-id='+node.replace(':','-')+'&starting-point-node-id='+node.replace(':','-')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(24);doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 2'].font.size=Pt(12.5)
for name in ['Heading 1','Heading 2']:doc.styles[name].paragraph_format.keep_with_next=True
doc.core_properties.title='Kịch bản thuyết trình và demo Power9 Office Routine';doc.core_properties.author='Nhóm Power9'
h=sec.header.paragraphs[0];h.text='Power9 Office Routine   PRM323 Lab 2';h.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('Power9 Office   ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def clean(t):
 t=re.sub(r'\b((?:OI|F|S|C|P)\d+)(?=[^\W\d_])',r'\1 ',t)
 protected={}
 def hold(m):
  key='⟦'+chr(65+len(protected)//26)+chr(65+len(protected)%26)+'⟧';protected[key]=m.group();return key
 t=re.sub(r'https?://\S+|MTv5qJbn065VYJQIXausoq|Power9|#[0-9a-fA-F]{6}|\b(?:OI|F|S|C|P)\d+\b',hold,t)
 t=re.sub(r'(?<=[^\W\d_])(?=\d)|(?<=\d)(?=[^\W\d_])',' ',t)
 t=re.sub(r'([,;])(?=\S)',r'\1 ',t)
 t=re.sub(r'(?<=[^\W\d_]):(?=\d)',': ',t)
 t=re.sub(r'\s*×\s*',' × ',t)
 for key,value in protected.items():t=t.replace(key,value)
 return t
def p(t,bold=False):
 a=doc.add_paragraph();a.add_run(clean(t)).bold=bold
 if a._p.getprevious() is not None and a._p.getprevious().tag==qn('w:tbl'):a.paragraph_format.space_before=Pt(6)
 return a
def link(par,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1E5AA8');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=clean(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=clean(v)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
   pr.append(border);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1E5AA8' if i==0 else 'FFFFFF' if i%2 else 'F3F6FA');pr.append(sh);mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.04
    for r in para.runs:r.font.name='Arial';r.font.size=Pt(9.5);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
 return t
def heading(t,level=1):doc.add_heading(clean(t),level)
def newpage(t):doc.add_page_break();heading(t)
def refs(sl):
 a=doc.add_paragraph();a.paragraph_format.space_after=Pt(9);a.add_run('Nguồn: ').italic=True
 for i,(label,url) in enumerate(sl['sources']):
  if i:a.add_run('  ')
  link(a,label,url)
 for r in a.runs:r.font.size=Pt(8.5)
timings=[35,55,45,55,50,75,50,45,55,50,40,45]
def clock(s):return str(s//60)+':'+str(s%60).zfill(2)
doc.add_paragraph('Power9 Office Routine','Title');doc.add_paragraph('Kịch bản thuyết trình và demo có giải thích','Subtitle')
p('Bản này đi cùng 12 slide Office mới. Lời nói phục vụ trình bày 10 phút và demo trực tiếp khoảng 5 phút. Thời gian là phân bổ dự kiến, chưa phải kết quả đo khi chạy với người dùng.')
p('Thiết kế có 10 màn chức năng, 3 starting point và accordion đủ 9 nhóm. Đã kiểm tra cấu trúc và dữ liệu; chưa có user test, TalkBack hoặc app Flutter. Timer, tìm kiếm và nhập giờ là các trạng thái mẫu của prototype.')
rows=[];elapsed=0
for sl,seconds in zip(slides,timings):rows.append([str(sl['n']).zfill(2),sl['title'],clock(elapsed)+'–'+clock(elapsed+seconds)]);elapsed+=seconds
table(['Slide','Nội dung','Khoảng trình bày'],rows,[1.3,12,4.6])
a=p('Mở sẵn: ');link(a,'Figma hiện hành',fig);a.add_run('   ');link(a,'Stitch dự án',stitch);a.add_run('   ');link(a,'GitHub','https://github.com/PhamQuangDuc1/PRM.POWER9')
p('Giữ PPTX, DOCX và sources cạnh nhau sau khi giải nén. Nguồn cục bộ của bản này nằm trong gói, không tự coi các thay đổi đã được đẩy lên GitHub.',True)
gestures={1:'Chỉ tên app, đối tượng và link nộp.',2:'Chỉ ba persona và nhấn mạnh mức ưu tiên ngang nhau.',3:'Chỉ bốn tab rồi sáu màn con.',4:'Nêu mục tiêu, đích và nhánh phục hồi.',5:'Chỉ prompt và output; mở nguồn nếu được hỏi.',6:'So sánh từng cặp trước/sau và nói rõ xuất xứ.',7:'Giải thích ít nhất một accept, modify và reject.',8:'Đối chiếu S07 và chỉ đủ mười màn.',9:'Chỉ token, instance và component.',10:'Đọc một tỷ lệ rồi chỉ 360/412.',11:'Giới thiệu ba link bắt đầu demo.',12:'Chỉ handoff và nêu đúng giới hạn.'}
for start in [0,4,8]:
 newpage('Lời thuyết trình slide '+str(start+1)+' đến '+str(start+4))
 for sl,seconds in zip(slides[start:start+4],timings[start:start+4]):
  heading('Slide '+str(sl['n']).zfill(2)+' '+sl['title']+' '+str(seconds)+' giây',2);p(narration[str(sl['n'])]);p('Trên màn hình: '+gestures[sl['n']]);refs(sl)
newpage('Demo chuẩn bị và flow F01')
p('Trước buổi chấm, mở thử quyền xem Figma, Stitch và GitHub bằng cửa sổ chưa đăng nhập. Mở đúng file MTv5qJbn065VYJQIXausoq, không dùng Family Routine cũ. Tắt thông báo cá nhân và mở sẵn ba starting point. Truy cập qua MCP chưa thay cho kiểm tra quyền công khai.')
table(['Mốc','Nội dung'],[['0:00–0:25','Cho xem 6 page và 10 màn Final UI'],['0:25–1:35','F01 Mai tìm/thêm, Back, lỗi và Thử lại'],['1:35–2:50','F02 Huy paused, overlay, kết thúc sớm, ghi nhận'],['2:50–3:45','F03 Linh đổi giờ và giờ yên lặng'],['3:45–4:25','Auto Layout, variants, instance, token binding'],['4:25–5:00','AI log, before/after và handoff']],[3.3,14.6])
p('Lời mở: “Đây là sáu page theo yêu cầu. Page 03 có mười màn chức năng. Các frame khác là dữ liệu và trạng thái phục vụ ba nhiệm vụ, không tính thêm màn.”')
a=p('F01 ');link(a,'Mai tìm và thêm',proto('30:2'))
table(['Thao tác','Lời giải thích','Kết quả cần thấy'],[
['Home: cuộn nội dung, mở Khám phá','“Mai có khoảng nghỉ ngắn nên lọc theo thời lượng.”','Khám phá, đủ 9 nhóm'],
['Mở/thu gọn nhóm khác 07, chọn ≤5 phút','“Mọi nhóm đều mở được. Bộ lọc áp dụng lên hoạt động.”','Thẻ có thời lượng; ẩn việc 10 phút'],
['Cuộn đến 07; chọn Rủ đồng nghiệp đi bộ','“Đây là hoạt động tự nguyện, app không tự gửi lời mời.”','Đúng tên, 5 phút'],
['Back, mở lại chi tiết, chọn lịch','“Back giữ bộ lọc, Mai không phải chọn lại.”','Vẫn ≤5 phút; ngày26/10,12:00'],
['Lưu, đợi lỗi; cuộn đến Thử lại','“Một lần lưu lỗi được mô phỏng. Tên và ngày giờ vẫn còn.”','Form không bị xóa dữ liệu'],
['Thử lại, mở Tiến độ sau thành công','“Lưu lịch không đánh dấu đã hoàn thành.”','Home1/4; Tiến độ25%']],[4.3,8.2,5.4])
p('Nhóm 07 mặc định đang mở, không cần chạm tiêu đề nếu danh sách đã hiện. Không nói có nhập liệu hoặc time picker thật: prototype chỉ có lựa chọn mẫu. Cuộn vùng điện thoại, không kéo canvas Figma.')
newpage('Demo flow F02 và kiểm soát phiên nghỉ')
a=p('F02 ');link(a,'Huy thực hiện và ghi nhận',proto('30:324'))
p('Mục tiêu là chứng minh quyền kiểm soát khi bị gián đoạn. Dùng nhánh kết thúc sớm để không chờ đủ ba phút. Mở starting point mới để về bộ dữ liệu gốc của flow.')
table(['Thao tác','Lời giải thích','Kết quả cần thấy'],[
['Mở Huy, tab Thư giãn','“Huy cần khoảng nghỉ ngắn sau cuộc họp.”','Chưa bắt đầu,03:00'],
['Bắt đầu','“Bản mẫu chuyển đến snapshot đã trôi45giây, còn02:15, không đếm từng giây.”','Đang thực hiện,02:15'],
['Tạm dừng','“Huy chủ động dừng khi có việc xen vào.”','Paused, vẫn02:15'],
['Rời phiên nghỉ','“Xác nhận giúp tránh mất phiên do chạm nhầm.”','Overlay trên nền cũ'],
['Tiếp tục trong dialog','“Nút này đóng dialog về đúng trạng thái trước.”','Vẫn paused, chưa ghi nhận'],
['Từ paused chọn Kết thúc','“Kết thúc sớm chỉ ghi thời gian đã thực hiện.”','Thông báo45giây'],
['Ghi nhận hoạt động','“Bản ghi hoàn thành tách khỏi lịch hoạt động.”','Tiến độ2/4;45giây'],
['Back về Home','“Kết quả vừa ghi nhận được giữ lại.”','Home2/4, phiên45giây']],[4.3,8.2,5.4])
p('Tiếp tục trong dialog chỉ đóng overlay về paused; để resume phải chọn Tiếp tục trên màn paused. App thực cần clock và xử lý pause/resume/background. Mẫu này chưa chứng minh phần đó.')
heading('Nhánh phụ khi có thêm thời gian',2)
p('Từ Home gốc, hoàn thành Đi bộ giữa giờ rồi Hoàn tác:2/3phục hồi1/3. Từ Tiến độ gốc, Xóa bản ghi mở overlay. Giữ lại không đổi dữ liệu; Xóa đưa Home/history về0/3nhưng giữ lịch. Không thay nhánh phụ này cho ba flow chính.')
newpage('Demo flow F03 và chứng minh bàn giao')
a=p('F03 ');link(a,'Linh đổi lịch',proto('30:638'))
table(['Thao tác','Lời giải thích','Kết quả cần thấy'],[
['Home: chọn Tổng kết công việc','“Linh sửa việc có sẵn khi lịch họp thay đổi.”','17:30; chi tiết đã thêm'],
['Đổi lịch','“Bản edit mô phỏng lựa chọn18:00, không thêm bản sao.”','Cùng hoạt động,18:00'],
['Chỉ phần giờ yên lặng','“Lịch vẫn lưu được, nhưng thông báo không phát trong giờ yên lặng.”','18:00–08:00'],
['Lưu, xem lỗi rồi Thử lại','“Lỗi giữ draft18:00; Home giữ17:30cho tới khi lưu thành công.”','Tên/ngày/giờ không đổi'],
['Xem Home và Tiến độ','“Đổi giờ không làm tăng tổng số hoạt động.”','18:00; vẫn1/3']],[4.3,8.2,5.4])
heading('40 giây chứng minh design system',2)
p('Chuyển về Design, page05, chọn P9/Button. Chỉ Auto Layout, states và instance. Trong fill/padding chỉ binding color/primary-container và spacing/16. Mở DiscoverGroup để thấy18variants. Page04có Variables/styles. Không cần sửa token thật trong buổi chấm.')
heading('35 giây chứng minh AI log và handoff',2)
p('Mở nguồn slide06, chọn cặp font trước/sau. Nói rõ đó là kiểm chứng Chrome cục bộ rồi mở phản biện C01để chỉ vấn đề, quyết định và bằng chứng. Mở sources/handoff/flutter-handoff.md, chỉ sáu mục của một màn và hai mapping. Kết: “Prototype mô phỏng hành vi; lập trình viên còn phải triển khai dữ liệu và thời gian thật.”')
p('Nếu demo chậm, mở lại đúng starting point từ slide11. Nếu mất mạng, trình bày ảnh trong sources và nói rõ đây là bằng chứng tĩnh, không nhận là demo tương tác.')
newpage('Câu hỏi và câu trả lời gợi ý')
qa=[
('Vì sao ba persona đều là chính?','Phạm vi đã chốt ưu tiên Huy, Mai, Linh ngang nhau. Họ có nhu cầu khác nhau trong cùng app; mỗi flow đại diện một nhiệm vụ, không có màn chọn persona.'),
('Vì sao không bắt hoàn thành chín nhóm mỗi ngày?','Chín nhóm là danh mục. Tiến độ tính trên hoạt động tự chọn, không phải chỉ tiêu9/9hoặc điểm sức khỏe.'),
('Nhóm từ chối đề xuất AI nào?','Bỏ nhịp tim ước tính và điểm sức khỏe vì không có dữ liệu đo hoặc tính năng tương ứng. Giữ nội dung phục vụ nhiệm vụ; xem OI03và slide07.'),
('Ba vòng có khác nhau không?','Vòng1sửa font; vòng2giảm mật độ; vòng3giữ dữ liệu và phục hồi lỗi. Mỗi vòng có ảnh trước/sau, ghi rõ Stitch, render cục bộ hoặc Figma.'),
('Vì sao nhiều frame và reaction hơn màn?','Chỉ có10màn chức năng. Frame khác là dữ liệu, bộ lọc, error, overlay, snapshot;323reactions còn gồm accordion, không cộng chúng thành số màn.'),
('Chứng minh token được dùng thế nào?','Chọn instance, truy về master rồi xem binding fill, padding, gap hoặc radius. Ảnh bảng màu không đủ chứng minh tái sử dụng.'),
('Đã đạt toàn bộ accessibility chưa?','Chưa. Đã tính các cặp màu, kiểm tra kích thước và layout360/412. Chưa thử TalkBack, text scaling hoặc người dùng, không tuyên bố chứng nhận WCAG.'),
('Người khác dựng từ handoff được chưa?','Có10màn×6mục và hai mapping, nhưng chưa có peer handoff test. Cần người khác dựng một màn để tìm chỗ còn phải hỏi.'),
('Prompt có đúng bản thực gửi không?','Gói giữ nguyên văn bản lưu và các prompt phản biện. Chưa đủ chat gửi Stitch để đối chiếu100%, nên không tự xác nhận thay lịch sử. Link Stitch mới giúp người trình bày mở dự án để kiểm tra.'),
('Link công khai đã kiểm tra chưa?','MCP truy cập được chưa thay cho cửa sổ chưa đăng nhập. Người trình bày cần thử quyền xem trước buổi chấm.')]
for q,a in qa:heading(q,2);p(a)
p('Nguồn yêu cầu: hai tài liệu hướng dẫn PRM323 Lab2, mục2–4. Nguồn sản phẩm: Figma hiện hành, checkpoint, audit accordion, handoff và AI log trong sources. Không cần đọc hết nguồn trong bài nói; mở khi có câu hỏi.')
for b in list(doc.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
for b in list(doc.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.save(out/'Power9_Office_Kich_Ban_Thuyet_Trinh_Va_Demo.docx');print('DOCX written')
